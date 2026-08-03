from __future__ import annotations

import csv
import json
import re
import sqlite3
from pathlib import Path

from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches, Pt

from app.config import settings
from app.models import AssessmentResult, ExportFormat, ExportMetadata, Finding, Recommendation


SENSITIVE_KEY_PATTERN = re.compile(r"(token|secret|password|authorization|cookie|connectionstring|client_secret)", re.IGNORECASE)
SENSITIVE_VALUE_PATTERN = re.compile(r"(bearer\s+[a-z0-9\-\._~\+\/]+=*|accountkey=|sig=)", re.IGNORECASE)
PPTX_SLIDE_TITLES = [
    "Executive Summary",
    "Assessment Scope and Methodology",
    "Security Overview",
    "Defender for Cloud Posture",
    "Azure Policy and Compliance",
    "Identity and Access Posture",
    "Management Group and Subscription Governance",
    "Network, Storage, and Key Vault Security",
    "Compute, AKS, and Logging Posture",
    "Security Recommendations Register",
]
PDF_SECTION_TITLES = [
    "Executive Summary",
    "Assessment Scope and Methodology",
    "Security Overview",
    "Findings by Severity",
    "Findings by Domain",
    "Partial Failures and Data Quality",
    "Recommendations Register",
    "Detailed Findings and Evidence",
]


def _sanitize(value: object, key: str = "") -> object:
    if SENSITIVE_KEY_PATTERN.search(key):
        return "[REDACTED]"
    if isinstance(value, dict):
        return {inner_key: _sanitize(inner_value, inner_key) for inner_key, inner_value in value.items()}
    if isinstance(value, list):
        return [_sanitize(item, key) for item in value]
    if isinstance(value, str) and SENSITIVE_VALUE_PATTERN.search(value):
        return "[REDACTED]"
    return value


def _export_dir() -> Path:
    directory = settings.exports_dir
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _flatten_finding(run_id: str, finding: Finding) -> dict[str, object]:
    return {
        "record_type": "finding",
        "run_id": run_id,
        "finding_id": finding.finding_id,
        "tenant_id": finding.tenant_id,
        "tenant_name": finding.tenant_name,
        "subscription_id": finding.subscription_id,
        "subscription_name": finding.subscription_name,
        "domain": finding.domain,
        "severity": finding.severity,
        "title": finding.title,
        "description": finding.description,
        "affected_resource_id": finding.affected_resource_id,
        "affected_resource_name": finding.affected_resource_name,
        "affected_resource_type": finding.affected_resource_type,
        "resource_group": finding.resource_group,
        "region": finding.region,
        "source_system": finding.source_system,
        "source_api": finding.source_api,
        "evidence_id": finding.evidence_id,
        "observed_at_utc": finding.observed_at_utc.isoformat(),
        "recommendation": finding.recommendation,
        "remediation_text": finding.remediation_text,
        "remediation_command_text": finding.remediation_command_text,
        "confidence": finding.confidence,
        "data_quality_notes": json.dumps(_sanitize(finding.data_quality_notes), ensure_ascii=True),
        "evidence": json.dumps(_sanitize([item.model_dump(by_alias=True, mode="json") for item in finding.evidence]), ensure_ascii=True),
    }


def _flatten_recommendation(run_id: str, recommendation: Recommendation) -> dict[str, object]:
    return {
        "record_type": "recommendation",
        "run_id": run_id,
        "recommendation_id": recommendation.recommendation_id,
        "finding_id": recommendation.finding_id,
        "title": recommendation.title,
        "priority": recommendation.priority,
        "recommendation": recommendation.recommendation,
        "remediation_text": recommendation.remediation_text,
        "remediation_command_text": recommendation.remediation_command_text,
        "source_system": recommendation.source_system,
        "source_api": recommendation.source_api,
        "evidence_id": recommendation.evidence_id,
        "confidence": recommendation.confidence,
        "data_quality_notes": json.dumps(_sanitize(recommendation.data_quality_notes), ensure_ascii=True),
    }


def export_csv(assessment: AssessmentResult) -> Path:
    output = _export_dir() / f"assessment_{assessment.metadata.run_id}.csv"
    rows = [_flatten_finding(assessment.metadata.run_id, finding) for finding in assessment.findings]
    rows.extend(_flatten_recommendation(assessment.metadata.run_id, recommendation) for recommendation in assessment.recommendations)

    headers = sorted({key for row in rows for key in row.keys()}) if rows else ["assessment_id", "finding_id"]
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _sanitize(value, key) for key, value in row.items()})
    return output


def export_pdf(assessment: AssessmentResult) -> Path:
    output = _export_dir() / f"assessment_{assessment.metadata.run_id}.pdf"
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Azure Security Posture Assessment", ln=True)
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 8, f"Run ID: {assessment.metadata.run_id}", ln=True)
    pdf.cell(0, 8, f"Tenant ID: {assessment.metadata.tenant_id}", ln=True)
    pdf.cell(0, 8, f"Subscriptions: {', '.join(assessment.metadata.subscription_ids)}", ln=True)
    pdf.ln(3)

    def section(title: str) -> None:
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, title, ln=True)
        pdf.set_font("Helvetica", size=9)

    section(PDF_SECTION_TITLES[0])
    pdf.multi_cell(0, 5, f"Grounded findings: {assessment.summary.total_findings}")
    pdf.multi_cell(0, 5, f"Recommendations: {assessment.summary.total_recommendations}")
    pdf.multi_cell(0, 5, f"Partial failures: {assessment.summary.partial_failure_count}")

    section(PDF_SECTION_TITLES[1])
    pdf.multi_cell(0, 5, "This report uses only live delegated Azure reads, validates tenant and subscription scope, and never executes remediation.")

    section(PDF_SECTION_TITLES[2])
    overview_findings = [item for item in assessment.findings if item.domain == "security_overview"]
    pdf.multi_cell(0, 5, f"Security overview finding count: {len(overview_findings)}")

    section(PDF_SECTION_TITLES[3])
    for severity, count in assessment.summary.findings_by_severity.items():
        pdf.multi_cell(0, 5, f"{severity}: {count}")

    section(PDF_SECTION_TITLES[4])
    for domain, count in assessment.summary.findings_by_domain.items():
        pdf.multi_cell(0, 5, f"{domain}: {count}")

    section(PDF_SECTION_TITLES[5])
    if not assessment.partial_failures:
        pdf.multi_cell(0, 5, "No partial failures were recorded.")
    for failure in assessment.partial_failures:
        pdf.multi_cell(0, 5, f"{failure.domain}: {failure.message}")

    section(PDF_SECTION_TITLES[6])
    if not assessment.recommendations:
        pdf.multi_cell(0, 5, "No recommendations were generated.")
    for recommendation in assessment.recommendations:
        pdf.multi_cell(0, 5, f"[{recommendation.priority}] {recommendation.title}: {recommendation.recommendation}")

    section(PDF_SECTION_TITLES[7])
    if not assessment.findings:
        pdf.multi_cell(0, 5, "No findings were returned by the source APIs.")
    for finding in assessment.findings:
        pdf.multi_cell(0, 5, f"[{finding.severity}] {finding.title}: {finding.description}")
        for evidence in finding.evidence:
            pdf.multi_cell(0, 5, f"- {evidence.summary}: {json.dumps(_sanitize(evidence.details), ensure_ascii=True)}")
        pdf.ln(1)
    pdf.output(str(output))
    return output


def export_pptx(assessment: AssessmentResult) -> Path:
    output = _export_dir() / f"assessment_{assessment.metadata.run_id}.pptx"
    presentation = Presentation()
    for title in PPTX_SLIDE_TITLES:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        text_frame = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(9), Inches(5)).text_frame
        text_frame.word_wrap = True
        if title == "Executive Summary":
            text_frame.text = f"Run ID: {assessment.metadata.run_id}"
            paragraph = text_frame.add_paragraph()
            paragraph.text = f"Tenant: {assessment.metadata.tenant_id}"
            paragraph.font.size = Pt(18)
            paragraph = text_frame.add_paragraph()
            paragraph.text = f"Findings: {assessment.summary.total_findings}; Recommendations: {assessment.summary.total_recommendations}; Partial failures: {assessment.summary.partial_failure_count}"
            paragraph.font.size = Pt(18)
            continue
        if title == "Assessment Scope and Methodology":
            text_frame.text = f"Subscriptions: {', '.join(assessment.metadata.subscription_ids)}"
            paragraph = text_frame.add_paragraph()
            paragraph.text = "Strict validation is applied and all findings are grounded in live read-only Azure responses."
            paragraph.font.size = Pt(18)
            continue
        title_domain_map = {
            "Security Overview": ["security_overview"],
            "Defender for Cloud Posture": ["defender_for_cloud_posture"],
            "Azure Policy and Compliance": ["azure_policy_and_compliance"],
            "Identity and Access Posture": ["identity_and_access_posture"],
            "Management Group and Subscription Governance": ["management_group_and_subscription_governance"],
            "Network, Storage, and Key Vault Security": ["network_security", "storage_security", "key_vault_security"],
            "Compute, AKS, and Logging Posture": ["compute_and_vm_security", "aks_and_container_security", "logging_monitoring_and_evidence_retention"],
        }
        if title == "Security Recommendations Register":
            items = assessment.recommendations[:10]
            if not items:
                text_frame.text = "No recommendations were generated."
            for item in items:
                paragraph = text_frame.add_paragraph()
                paragraph.text = f"[{item.priority}] {item.title}"
                paragraph.font.size = Pt(18)
            continue
        findings = [finding for finding in assessment.findings if finding.domain in title_domain_map.get(title, [])]
        if not findings:
            text_frame.text = "No findings returned for this topic."
            continue
        for finding in findings[:8]:
            paragraph = text_frame.add_paragraph()
            paragraph.text = f"[{finding.severity}] {finding.title}"
            paragraph.font.size = Pt(18)
            detail = text_frame.add_paragraph()
            detail.text = finding.description
            detail.level = 1
            detail.font.size = Pt(12)

    presentation.save(output)
    return output


def export_sqlite(assessment: AssessmentResult) -> Path:
    output = _export_dir() / f"assessment_{assessment.metadata.run_id}.sqlite"
    if output.exists():
        output.unlink()
    connection = sqlite3.connect(output)
    cursor = connection.cursor()

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS assessment_runs (
            run_id TEXT PRIMARY KEY,
            tenant_id TEXT NOT NULL,
            tenant_name TEXT,
            strict_validation INTEGER NOT NULL,
            grounded INTEGER NOT NULL,
            started_at_utc TEXT NOT NULL,
            completed_at_utc TEXT NOT NULL,
            finding_count INTEGER NOT NULL,
            recommendation_count INTEGER NOT NULL,
            partial_failure_count INTEGER NOT NULL
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS tenants (
            tenant_id TEXT PRIMARY KEY,
            tenant_name TEXT
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS subscriptions (
            subscription_id TEXT PRIMARY KEY,
            tenant_id TEXT NOT NULL,
            subscription_name TEXT,
            state TEXT
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS resources (
            resource_id TEXT PRIMARY KEY,
            subscription_id TEXT NOT NULL,
            resource_name TEXT,
            resource_type TEXT,
            resource_group TEXT,
            region TEXT
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS findings (
            finding_id TEXT PRIMARY KEY,
            run_id TEXT NOT NULL,
            tenant_id TEXT NOT NULL,
            tenant_name TEXT,
            subscription_id TEXT NOT NULL,
            subscription_name TEXT,
            domain TEXT NOT NULL,
            severity TEXT NOT NULL,
            title TEXT NOT NULL,
            description TEXT NOT NULL,
            affected_resource_id TEXT NOT NULL,
            affected_resource_name TEXT,
            affected_resource_type TEXT,
            resource_group TEXT,
            region TEXT,
            source_system TEXT NOT NULL,
            source_api TEXT NOT NULL,
            evidence_id TEXT NOT NULL,
            observed_at_utc TEXT NOT NULL,
            recommendation TEXT NOT NULL,
            remediation_text TEXT NOT NULL,
            remediation_command_text TEXT,
            confidence TEXT NOT NULL,
            data_quality_notes_json TEXT NOT NULL
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS recommendations (
            recommendation_id TEXT PRIMARY KEY,
            run_id TEXT NOT NULL,
            finding_id TEXT NOT NULL,
            title TEXT NOT NULL,
            priority TEXT NOT NULL,
            recommendation TEXT NOT NULL,
            remediation_text TEXT NOT NULL,
            remediation_command_text TEXT,
            source_system TEXT NOT NULL,
            source_api TEXT NOT NULL,
            evidence_id TEXT NOT NULL,
            confidence TEXT NOT NULL,
            data_quality_notes_json TEXT NOT NULL
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS evidence (
            evidence_id TEXT PRIMARY KEY,
            finding_id TEXT NOT NULL,
            run_id TEXT NOT NULL,
            source_system TEXT NOT NULL,
            source_api TEXT NOT NULL,
            observed_at_utc TEXT NOT NULL,
            record_type TEXT NOT NULL,
            summary TEXT NOT NULL,
            details_json TEXT NOT NULL,
            data_quality_notes_json TEXT NOT NULL
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS api_query_log (
            query_id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id TEXT NOT NULL,
            category TEXT NOT NULL,
            api_route TEXT NOT NULL,
            tenant_id TEXT,
            subscription_id TEXT,
            source_system TEXT,
            source_api TEXT NOT NULL,
            endpoint TEXT NOT NULL,
            query_text TEXT,
            status_code INTEGER,
            message TEXT,
            observed_at_utc TEXT NOT NULL,
            remediation_guidance TEXT
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS export_artifacts (
            export_id TEXT PRIMARY KEY,
            run_id TEXT NOT NULL,
            format TEXT NOT NULL,
            file_name TEXT NOT NULL,
            generated_at_utc TEXT NOT NULL,
            record_count INTEGER NOT NULL
        )
        """
    )

    cursor.execute(
        "INSERT INTO assessment_runs VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            assessment.metadata.run_id,
            assessment.metadata.tenant_id,
            assessment.metadata.tenant_name,
            1 if assessment.metadata.strict_validation else 0,
            1 if assessment.metadata.grounded else 0,
            assessment.metadata.started_at_utc.isoformat(),
            assessment.metadata.completed_at_utc.isoformat() if assessment.metadata.completed_at_utc else "",
            assessment.metadata.finding_count,
            assessment.metadata.recommendation_count,
            assessment.metadata.partial_failure_count,
        ),
    )

    cursor.execute("INSERT INTO tenants VALUES (?, ?)", (assessment.metadata.tenant_id, assessment.metadata.tenant_name))
    for subscription_id, subscription_name in zip(assessment.metadata.subscription_ids, assessment.metadata.subscription_names, strict=False):
        cursor.execute("INSERT INTO subscriptions VALUES (?, ?, ?, ?)", (subscription_id, assessment.metadata.tenant_id, subscription_name, "Selected"))

    inserted_resources: set[str] = set()
    for finding in assessment.findings:
        if finding.affected_resource_id not in inserted_resources:
            inserted_resources.add(finding.affected_resource_id)
            cursor.execute(
                "INSERT INTO resources VALUES (?, ?, ?, ?, ?, ?)",
                (
                    finding.affected_resource_id,
                    finding.subscription_id,
                    finding.affected_resource_name,
                    finding.affected_resource_type,
                    finding.resource_group,
                    finding.region,
                ),
            )
        cursor.execute(
            "INSERT INTO findings VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                finding.finding_id,
                assessment.metadata.run_id,
                finding.tenant_id,
                finding.tenant_name,
                finding.subscription_id,
                finding.subscription_name,
                finding.domain,
                finding.severity,
                finding.title,
                finding.description,
                finding.affected_resource_id,
                finding.affected_resource_name,
                finding.affected_resource_type,
                finding.resource_group,
                finding.region,
                finding.source_system,
                finding.source_api,
                finding.evidence_id,
                finding.observed_at_utc.isoformat(),
                finding.recommendation,
                finding.remediation_text,
                finding.remediation_command_text,
                finding.confidence,
                json.dumps(_sanitize(finding.data_quality_notes), ensure_ascii=True),
            ),
        )
        for item in finding.evidence:
            cursor.execute(
                "INSERT INTO evidence VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    item.evidence_id,
                    finding.finding_id,
                    assessment.metadata.run_id,
                    item.source_system,
                    item.source_api,
                    item.observed_at_utc.isoformat(),
                    item.record_type,
                    item.summary,
                    json.dumps(_sanitize(item.details), ensure_ascii=True),
                    json.dumps(_sanitize(item.data_quality_notes), ensure_ascii=True),
                ),
            )
            cursor.execute(
                "INSERT INTO api_query_log (run_id, category, api_route, tenant_id, subscription_id, source_system, source_api, endpoint, query_text, status_code, message, observed_at_utc, remediation_guidance) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    assessment.metadata.run_id,
                    "success",
                    "/api/assessment/run",
                    finding.tenant_id,
                    finding.subscription_id,
                    item.source_system,
                    item.source_api,
                    str(item.details.get("endpoint") or ""),
                    str(item.details.get("query") or ""),
                    item.details.get("statusCode"),
                    item.summary,
                    item.observed_at_utc.isoformat(),
                    None,
                ),
            )

    for recommendation in assessment.recommendations:
        cursor.execute(
            "INSERT INTO recommendations VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                recommendation.recommendation_id,
                assessment.metadata.run_id,
                recommendation.finding_id,
                recommendation.title,
                recommendation.priority,
                recommendation.recommendation,
                recommendation.remediation_text,
                recommendation.remediation_command_text,
                recommendation.source_system,
                recommendation.source_api,
                recommendation.evidence_id,
                recommendation.confidence,
                json.dumps(_sanitize(recommendation.data_quality_notes), ensure_ascii=True),
            ),
        )

    for failure in assessment.partial_failures:
        cursor.execute(
            "INSERT INTO api_query_log (run_id, category, api_route, tenant_id, subscription_id, source_system, source_api, endpoint, query_text, status_code, message, observed_at_utc, remediation_guidance) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                assessment.metadata.run_id,
                failure.category,
                failure.api_route,
                failure.tenant_id,
                failure.subscription_id,
                failure.source_system,
                failure.source_api,
                "",
                "",
                failure.status_code,
                failure.message,
                failure.observed_at_utc.isoformat(),
                failure.remediation_guidance,
            ),
        )

    export_metadata = ExportMetadata(
        export_id=f"sqlite-{assessment.metadata.run_id}",
        run_id=assessment.metadata.run_id,
        format="sqlite",
        file_name=output.name,
        generated_at_utc=assessment.metadata.completed_at_utc or assessment.metadata.started_at_utc,
        record_count=len(assessment.findings) + len(assessment.recommendations),
    )
    cursor.execute(
        "INSERT INTO export_artifacts VALUES (?, ?, ?, ?, ?, ?)",
        (
            export_metadata.export_id,
            export_metadata.run_id,
            export_metadata.format,
            export_metadata.file_name,
            export_metadata.generated_at_utc.isoformat(),
            export_metadata.record_count,
        ),
    )

    connection.commit()
    connection.close()
    return output


EXPORTERS = {
    ExportFormat.csv: export_csv,
    ExportFormat.pdf: export_pdf,
    ExportFormat.pptx: export_pptx,
    ExportFormat.sqlite: export_sqlite,
}


def export_assessment(assessment: AssessmentResult, export_format: ExportFormat) -> Path:
    return EXPORTERS[export_format](assessment)
